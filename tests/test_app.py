import io
import os
from datetime import datetime

from openpyxl import load_workbook

from mint import create_app

TEMPLATE = os.path.join(os.path.dirname(os.path.dirname(__file__)), "data",
                        "work_order_template.xlsx")


def make_wo(c, **extra):
    form = {"wo_number": "WO-1", "client": "Acme", "wo_status": "In progress",
            "direct_cost": "1,000", "target_markup": "40", "invoice_revenue": "$1,400.00",
            "overhead": "100", "admin_fees": "30", "other_expenses": "20",
            "approved_nte": "1500", "amount_received": "0",
            "invoice_sent_date": "2026-01-10", "payment_terms": "30",
            "invoice_status": "Sent"}
    form.update(extra)
    return c.post("/work-orders/new", data=form)


def filled_template(rows):
    wb = load_workbook(TEMPLATE)
    ws = wb["Work Orders"]
    for i, row in enumerate(rows):
        for col, v in row.items():
            ws[f"{col}{11 + i}"] = v
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def test_pages_render(client):
    for url in ("/", "/work-orders", "/work-orders/new", "/import", "/guide",
                "/work-orders?view=all", "/api/work-orders"):
        assert client.get(url).status_code == 200, url


def test_create_and_view_work_order(client):
    r = make_wo(client)
    assert r.status_code == 302
    page = client.get(r.headers["Location"]).get_data(as_text=True)
    assert "<h1>WO-1 " in page
    assert "$250.00" in page            # net profit
    assert "Invoice within NTE" in page
    data = client.get("/api/work-orders").get_json()
    wo = data["work_orders"][0]
    assert wo["target_markup"] == 0.4 and wo["quote"] == 1400.0
    assert data["totals"]["work_orders"] == 1


def test_validation_and_duplicates(client):
    r = make_wo(client, direct_cost="-5", wo_status="Bogus")
    assert r.status_code == 422
    body = r.get_data(as_text=True)
    assert "cannot be negative" in body and "is not one of" in body
    assert make_wo(client).status_code == 302
    r = make_wo(client)
    assert r.status_code == 422 and "already exists" in r.get_data(as_text=True)
    r = client.post("/work-orders/new", data={"wo_number": ""})
    assert r.status_code == 422


def test_edit_logs_changes_and_payment(client):
    loc = make_wo(client).headers["Location"]
    wo_id = int(loc.rsplit("/", 1)[1])
    page = client.get(f"/work-orders/{wo_id}/edit").get_data(as_text=True)
    assert 'value="40"' in page         # percent shown as typed
    client.post(f"/work-orders/{wo_id}/payment", data={"amount": "400", "method": "ACH"})
    wo = client.get("/api/work-orders").get_json()["work_orders"][0]
    assert wo["amount_received"] == 400 and wo["balance_due"] == 1000
    assert wo["invoice_status"] == "Partially paid" and wo["payment_method"] == "ACH"
    client.post(f"/work-orders/{wo_id}/payment", data={"amount": "1000"})
    wo = client.get("/api/work-orders").get_json()["work_orders"][0]
    assert wo["invoice_status"] == "Paid" and wo["balance_due"] == 0
    history = client.get(f"/work-orders/{wo_id}").get_data(as_text=True)
    assert "Payment of $1,000.00 recorded" in history


def test_filters_and_delete(client):
    make_wo(client)
    make_wo(client, wo_number="WO-2", client="Beta", invoice_revenue="", direct_cost="")
    page = client.get("/work-orders?flag=needs_input").get_data(as_text=True)
    assert "WO-2" in page and ">WO-1<" not in page
    page = client.get("/work-orders?q=beta").get_data(as_text=True)
    assert "WO-2" in page
    client.post("/work-orders/2/delete")
    assert len(client.get("/api/work-orders").get_json()["work_orders"]) == 1


def test_import_from_the_excel_template(client):
    buf = filled_template([
        {"A": "1001", "B": "Acme", "H": 1000, "I": 1500, "Y": 0.4, "Z": 1400,
         "AB": 100, "AC": 30, "AD": 20, "AG": 0, "R": "Completed",
         "U": datetime(2026, 1, 10), "AI": 30, "V": "Sent"},
        {"A": 1002, "B": "Beta", "R": "in progress"},          # case-insensitive choice
        {"A": "1003", "R": "Not a status"},                    # skipped
        {"B": "no wo number"},                                 # skipped
    ])
    r = client.post("/import", data={"file": (buf, "list.xlsx")},
                    content_type="multipart/form-data")
    body = r.get_data(as_text=True)
    assert r.status_code == 200
    wos = {w["wo_number"]: w for w in client.get("/api/work-orders").get_json()["work_orders"]}
    assert set(wos) == {"1001", "1002"}
    assert wos["1001"]["net_profit"] == 250 and wos["1001"]["payment_due"] == "2026-02-09"
    assert wos["1002"]["wo_status"] == "In progress"
    assert "Not a status" in body and "WO # is empty" in body

    # Re-import updates instead of duplicating.
    buf = filled_template([{"A": "1001", "Z": 1450}])
    client.post("/import", data={"file": (buf, "list.xlsx")}, content_type="multipart/form-data")
    wos = {w["wo_number"]: w for w in client.get("/api/work-orders").get_json()["work_orders"]}
    assert len(wos) == 2 and wos["1001"]["invoice_revenue"] == 1450
    assert wos["1001"]["client"] is None         # a blank cell in the sheet clears the value

    # Columns that are not in the file at all are left untouched.
    csv_file = io.BytesIO(b"WO #,Client,Invoice status\n1001,Acme Corp,paid\n")
    client.post("/import", data={"file": (csv_file, "x.csv")}, content_type="multipart/form-data")
    wos = {w["wo_number"]: w for w in client.get("/api/work-orders").get_json()["work_orders"]}
    assert wos["1001"]["client"] == "Acme Corp" and wos["1001"]["invoice_status"] == "Paid"
    assert wos["1001"]["invoice_revenue"] == 1450


def test_import_rejects_wrong_layout(client):
    buf = io.BytesIO(b"name,amount\nx,1\n")
    r = client.post("/import", data={"file": (buf, "x.csv")}, content_type="multipart/form-data")
    assert r.status_code == 400


def test_export_xlsx_keeps_template_formulas_and_round_trips(client, app, tmp_path):
    make_wo(client)
    r = client.get("/export.xlsx")
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.data))
    ws = wb["Work Orders"]
    assert ws["A11"].value == "WO-1" and ws["H11"].value == 1000 and ws["Y11"].value == 0.4
    assert ws["J11"].value.startswith("=IF(")
    assert ws["A5"].value == "=COUNTA(A11:A510)"
    assert "Guide and Example" in wb.sheetnames

    # The exported file imports cleanly into a fresh database.
    other = create_app({"DATABASE": str(tmp_path / "other.sqlite3"), "TESTING_NO_CSRF": True})
    c2 = other.test_client()
    c2.post("/import", data={"file": (io.BytesIO(r.data), "export.xlsx")},
            content_type="multipart/form-data")
    wo = c2.get("/api/work-orders").get_json()["work_orders"][0]
    assert wo["net_profit"] == 250 and wo["target_markup"] == 0.4


def test_export_grows_beyond_500_rows(app):
    from mint import excel
    rows = [{"wo_number": f"W{i}", "direct_cost": 1.0} for i in range(520)]
    wb = load_workbook(excel.export_xlsx(TEMPLATE, rows))
    ws = wb["Work Orders"]
    assert ws["A530"].value == "W519"
    assert ws["J530"].value == '=IF(OR(A530="",COUNT(H530)=0,COUNT(Z530)=0),"",ROUND(Z530-H530,2))'
    assert ws["A5"].value == "=COUNTA(A11:A530)"
    assert ws.tables["MintresponseWorkOrders"].ref == "A10:AK530"
    assert any("R11:R530" in str(dv.sqref) for dv in ws.data_validations.dataValidation)


def test_export_csv(client):
    make_wo(client)
    text = client.get("/export.csv").data.decode("utf-8-sig")
    header, row = text.splitlines()[:2]
    assert header.startswith("WO #,Client,Assignees")
    assert "40%" in row and "$250.00" in row


def test_csrf_and_password(tmp_path):
    app = create_app({"DATABASE": str(tmp_path / "s.sqlite3"), "APP_PASSWORD": "secret"})
    c = app.test_client()
    assert c.get("/").status_code == 302
    login = c.get("/login").get_data(as_text=True)
    token = login.split('name="csrf" value="')[1].split('"')[0]
    assert c.post("/login", data={"password": "secret"}).status_code == 400   # no CSRF token
    assert c.post("/login", data={"password": "wrong", "csrf": token}).status_code == 200
    assert c.post("/login", data={"password": "secret", "csrf": token}).status_code == 302
    assert c.get("/").status_code == 200
