"""The Excel application workbook (build_excel_app.py).

The structure tests always run. The recalculation test opens the workbook in
LibreOffice (when installed), lets it calculate every formula, and checks the
results against the web app's calculation engine.
"""
import datetime
import shutil
import subprocess

import pytest
from openpyxl import load_workbook

import build_excel_app as xa
from mint import calc
from seed_demo import demo_rows


@pytest.fixture(scope="module")
def demo_path(tmp_path_factory):
    path = tmp_path_factory.mktemp("xl") / "app.xlsx"
    xa.build(str(path), demo_rows())
    return path


def test_structure(demo_path):
    wb = load_workbook(demo_path)
    assert wb.sheetnames == ["Dashboard", "Work Orders", "Work Order Card", "Receivables",
                             "Clients", "Guide and Example"]
    ws = wb["Work Orders"]
    table = ws.tables["MintresponseWorkOrders"]
    assert table.ref == "A10:AP1010"
    assert [c.name for c in table.tableColumns][-5:] == [h[1] for h in xa.HELPERS]
    assert ws["J1010"].value.startswith('=IF(OR(A1010=""')
    assert ws["A5"].value == "=COUNTA(A11:A1010)"
    assert ws["A11"].value == "WO-24001" and isinstance(ws["T11"].value, datetime.datetime)
    for name in ("Dashboard", "Receivables", "Clients", "Work Order Card"):
        assert wb[name].protection.sheet
    assert not wb["Work Orders"].protection.sheet
    assert wb["Work Order Card"]["C5"].protection.locked is False
    assert len(wb["Dashboard"]._charts) == 3


def test_empty_workbook_builds(tmp_path):
    path = tmp_path / "empty.xlsx"
    xa.build(str(path))
    ws = load_workbook(path)["Work Orders"]
    assert ws["A11"].value is None and ws["AL11"].value.startswith("=IF(")


@pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice not installed")
def test_recalculated_values_match_the_app(demo_path, tmp_path):
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx",
                    "--outdir", str(tmp_path), str(demo_path)],
                   check=True, capture_output=True, timeout=300)
    wb = load_workbook(tmp_path / demo_path.name, data_only=True)
    d, rc, cl = wb["Dashboard"], wb["Receivables"], wb["Clients"]
    rows = [{**r, **calc.compute(r)} for r in demo_rows()]
    t = calc.totals(rows)

    tiles = {"B6": "work_orders", "E6": "invoiced_revenue", "H6": "gross_profit",
             "K6": "net_profit", "B10": "balance_due", "E10": "overdue_amount",
             "H10": "needs_input", "K10": "nte_exceptions"}
    for cell, key in tiles.items():
        assert d[cell].value == pytest.approx(t[key]), key

    for r in range(30, 36):
        bucket = d[f"B{r}"].value
        sel = [x for x in rows if x["aging_bucket"] == bucket]
        assert d[f"C{r}"].value == pytest.approx(sum(x["balance_due"] for x in sel)), bucket
        assert d[f"D{r}"].value == len(sel), bucket

    today = datetime.date.today()
    owed = sorted((x for x in rows if x["owed"]), key=lambda x: -(
        (today - datetime.date.fromisoformat(x["payment_due"])).days
        if x["payment_due"] else -100000))
    listed = [rc[f"C{r}"].value or "" for r in range(7, 7 + len(owed) + 1)]
    assert listed == [x["wo_number"] for x in owed] + [""]
    assert rc["D4"].value == pytest.approx(sum(x["balance_due"] for x in owed))
    assert rc["G4"].value == pytest.approx(t["overdue_amount"])

    revenue = {}
    for x in rows:
        revenue[x["client"]] = revenue.get(x["client"], 0) + (x.get("invoice_revenue") or 0)
    got = {cl[f"B{r}"].value: cl[f"D{r}"].value for r in range(6, 6 + len(revenue))}
    assert got == pytest.approx(revenue)
    top = sorted(((n, v) for n, v in revenue.items() if v > 0), key=lambda kv: -kv[1])[:8]
    assert [d[f"B{44 + i}"].value for i in range(len(top))] == [n for n, _ in top]
