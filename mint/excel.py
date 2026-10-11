"""Excel / CSV import and export in the Mintresponse Work Order List layout."""
import csv
import io
import re
from copy import copy
from datetime import date

from openpyxl import load_workbook
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.formula.translate import Translator
from openpyxl.worksheet.cell_range import MultiCellRange

from .fields import BY_LABEL, FIELDS, INPUT_FIELDS
from .values import InvalidValue, display, parse

SHEET = "Work Orders"
HEADER_ROW = 10
FIRST_ROW = 11
TEMPLATE_LAST_ROW = 510


# ---------------------------------------------------------------- import

def _normalise(label):
    return re.sub(r"[\s*]+", " ", str(label or "")).strip().lower()


def _find_header(ws, max_scan=40):
    """Locate the header row (the one whose cells include 'WO #')."""
    for row in ws.iter_rows(min_row=1, max_row=max_scan):
        labels = {c.column: _normalise(c.value) for c in row if c.value is not None}
        if "wo #" in labels.values():
            mapping = {col: BY_LABEL[lbl] for col, lbl in labels.items()
                       if lbl in BY_LABEL and not BY_LABEL[lbl].computed}
            return row[0].row, mapping
    return None, None


def read_rows(filename, stream):
    """Yield (source row number, {key: raw value}) from an .xlsx or .csv upload.

    Raises ValueError when no 'WO #' header can be found.
    """
    if filename.lower().endswith(".csv"):
        text = io.TextIOWrapper(stream, encoding="utf-8-sig", newline="")
        reader = csv.reader(text)
        header = next(reader, [])
        mapping = {i: BY_LABEL[_normalise(h)] for i, h in enumerate(header)
                   if _normalise(h) in BY_LABEL and not BY_LABEL[_normalise(h)].computed}
        if not any(f.key == "wo_number" for f in mapping.values()):
            raise ValueError("No 'WO #' column found in the CSV header.")
        for n, line in enumerate(reader, start=2):
            yield n, {f.key: line[i] if i < len(line) else None
                      for i, f in mapping.items()}
        return

    wb = load_workbook(stream, data_only=True, read_only=False)
    sheets = [wb[SHEET]] if SHEET in wb.sheetnames else []
    sheets += [ws for ws in wb.worksheets if ws.title != SHEET]
    for ws in sheets:
        header_row, mapping = _find_header(ws)
        if header_row:
            break
    else:
        raise ValueError("Could not find a header row with 'WO #'. "
                         "Use the Mintresponse Work Order List layout.")
    for row in ws.iter_rows(min_row=header_row + 1):
        values = {mapping[c.column].key: c.value for c in row if c.column in mapping}
        if any(v not in (None, "") for v in values.values()):
            yield row[0].row, values


def parse_row(raw):
    """Raw import values -> (clean data, list of error strings)."""
    data, errors = {}, []
    for key, value in raw.items():
        field = next(f for f in INPUT_FIELDS if f.key == key)
        try:
            data[key] = parse(field, value, source="excel")
        except InvalidValue as e:
            errors.append(str(e))
    return data, errors


# ---------------------------------------------------------------- export

def _extend_ref(ref, last):
    return re.sub(rf"(?<=\D){TEMPLATE_LAST_ROW}\b", str(last), ref)


def _grow_template(ws, last):
    """Extend formulas, styles, validation and summary ranges beyond row 510."""
    template = {c.column: c for c in ws[FIRST_ROW]}
    for r in range(TEMPLATE_LAST_ROW + 1, last + 1):
        for col, src in template.items():
            cell = ws.cell(row=r, column=col)
            cell._style = copy(src._style)
            if isinstance(src.value, str) and src.value.startswith("="):
                cell.value = Translator(src.value, origin=src.coordinate) \
                    .translate_formula(cell.coordinate)
    for c in ws[5]:
        if isinstance(c.value, str) and c.value.startswith("="):
            c.value = _extend_ref(c.value, last)
    for table in ws.tables.values():
        table.ref = _extend_ref(table.ref, last)
        if table.autoFilter is not None:
            table.autoFilter.ref = table.ref
    for dv in ws.data_validations.dataValidation:
        dv.sqref = MultiCellRange(_extend_ref(str(dv.sqref), last))
    old = ws.conditional_formatting
    ws.conditional_formatting = ConditionalFormattingList()
    for cf in old:
        for rule in cf.rules:
            ws.conditional_formatting.add(_extend_ref(str(cf.sqref), last), rule)


def export_xlsx(template_path, work_orders):
    """Fill the original tracker template so Excel formulas keep working."""
    wb = load_workbook(template_path)
    ws = wb[SHEET]
    last = max(TEMPLATE_LAST_ROW, FIRST_ROW + len(work_orders) - 1)
    if last > TEMPLATE_LAST_ROW:
        _grow_template(ws, last)
    for i, wo in enumerate(work_orders):
        r = FIRST_ROW + i
        for f in INPUT_FIELDS:
            v = wo.get(f.key)
            if v is not None and f.kind == "date":
                v = date.fromisoformat(v)
            ws[f"{f.col}{r}"].value = v
    wb.calculation.fullCalcOnLoad = True    # Excel recalculates every formula on open
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out


def export_csv(work_orders):
    """All 37 columns, formatted values, Excel column order."""
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow([f.label for f in FIELDS])
    for wo in work_orders:
        row = []
        for f in FIELDS:
            v = wo.get(f.key)
            if f.computed:
                v = display(f, v)
            elif f.kind == "percent" and v is not None:
                v = f"{v * 100:g}%"
            row.append("" if v is None else v)
        w.writerow(row)
    return io.BytesIO(out.getvalue().encode("utf-8-sig"))
