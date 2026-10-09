"""Build the Mintresponse Work Order *Excel application*.

A single .xlsx workbook that works on its own in Excel (no macros, no server):

  Dashboard          KPI tiles, charts, most-overdue invoices
  Work Orders        the original tracker table (yellow = enter, blue/green = formulas)
                     with 1,000 prepared rows and receivables helper columns
  Work Order Card    choose a WO # and see the whole job on one printable page
  Receivables        every open balance, most overdue first
  Clients            revenue, profit and balance per client
  Guide and Example  definitions from the original tracker

Usage:
    python build_excel_app.py                 -> Mintresponse_Work_Order_App.xlsx (empty)
    python build_excel_app.py --demo out.xlsx -> same workbook with 40 sample work orders
"""
import argparse
import os
from copy import copy
from datetime import date

from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.drawing.image import Image
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import TableColumn

from mint import excel
from mint.fields import BY_KEY, FORM_SECTIONS, INPUT_FIELDS, WO_STATUSES

HERE = os.path.dirname(os.path.abspath(__file__))
TEMPLATE = os.path.join(HERE, "data", "work_order_template.xlsx")
LOGO = os.path.join(HERE, "mint", "static", "logo.png")

FIRST, LAST = 11, 1010          # work-order rows on the Work Orders sheet
CAPACITY = LAST - FIRST + 1
WO = "'Work Orders'!"

NAVY, INK, MUTED = "203D56", "203040", "526578"
GREEN = "2F7D3A"
INPUT_FILL = PatternFill("solid", fgColor="FFF7DF")
CALC_FILL = PatternFill("solid", fgColor="EDF3F8")
TILE_FILL = PatternFill("solid", fgColor="F3F6F9")
WARN_FILL = PatternFill("solid", fgColor="FFF1D6")
HEAD_FILL = PatternFill("solid", fgColor=NAVY)
HEAD_FONT = Font(bold=True, color="FFFFFF", size=10)
THIN = Side(style="thin", color="D5DCE3")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
LINE = Border(bottom=THIN)

MONEY = '\\$#,##0.00;[Red]\\(\\$#,##0.00\\);\\$0.00'
PCT = "0.0%"
DATE = "mm/dd/yy"
AGING = ["Not yet due", "1-30 days", "31-60 days", "61-90 days", "Over 90 days",
         "No due date"]

# Receivables helper columns added to the right of the original table (AK).
HELPERS = [
    ("AL", "Aging bucket", 16, None),
    ("AM", "Days overdue", 14, "0"),
    ("AN", "AR sort key", 12, "0"),        # hidden
    ("AO", "AR order", 10, "0"),           # hidden
    ("AP", "Client #", 10, "0"),           # hidden
]


def rng(col):
    return f"{WO}${col}${FIRST}:${col}${LAST}"


def fmt_for(kind):
    return {"money": MONEY, "percent": PCT, "date": DATE, "int": "0"}.get(kind, "General")


def protect(ws):
    ws.protection.sheet = True
    ws.protection.formatColumns = False
    ws.protection.formatRows = False
    ws.protection.sort = False
    ws.protection.autoFilter = False


def header_row(ws, row, col, labels, widths=None):
    for i, label in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=label)
        c.fill, c.font = HEAD_FILL, HEAD_FONT
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30


def title(ws, text, subtitle=None, col="B"):
    ws[f"{col}2"] = text
    ws[f"{col}2"].font = Font(bold=True, size=16, color=NAVY)
    if subtitle:
        ws[f"{col}3"] = subtitle
        ws[f"{col}3"].font = Font(size=10, color=MUTED)
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 2


# ------------------------------------------------------------------ Work Orders

def build_work_orders(ws):
    excel._grow_template(ws, LAST)
    ws["A7"] = (f"Yellow cells = enter data. Blue/green cells = formulas. {CAPACITY:,} prepared "
                "rows. Dashboard, Receivables and Clients update automatically.")
    ws.freeze_panes = "B11"

    head_src, body_src = ws["AK10"], ws["AK11"]
    for col, label, width, numfmt in HELPERS:
        h = ws[f"{col}10"]
        h.value = label
        h._style = copy(head_src._style)
        ws.column_dimensions[col].width = width
    for r in range(FIRST, LAST + 1):
        f = {
            "AL": (f'=IF(A{r}="","",IF(NOT(AND(COUNT(AH{r})=1,N(AH{r})>0.004,V{r}<>"Void",'
                   f'V{r}<>"Not invoiced",V{r}<>"Draft")),"",IF(COUNT(X{r})=0,"No due date",'
                   f'IF(TODAY()-X{r}<=0,"Not yet due",IF(TODAY()-X{r}<=30,"1-30 days",'
                   f'IF(TODAY()-X{r}<=60,"31-60 days",IF(TODAY()-X{r}<=90,"61-90 days",'
                   f'"Over 90 days")))))))'),
            "AM": f'=IF(OR(AL{r}="",AL{r}="No due date"),"",MAX(TODAY()-X{r},0))',
            "AN": f'=IF(AL{r}="","",IF(COUNT(X{r})=0,-100000,TODAY()-X{r}))',
            "AO": (f'=IF(AN{r}="","",COUNTIF(AN${FIRST}:AN${LAST},">"&AN{r})'
                   f'+COUNTIF(AN${FIRST}:AN{r},AN{r}))'),
            "AP": (f'=IF(OR(A{r}="",B{r}=""),"",IF(COUNTIF(B${FIRST}:B{r},B{r})=1,'
                   f'MAX(AP${FIRST - 1}:AP{r - 1})+1,""))'),
        }
        for col, _label, _w, numfmt in HELPERS:
            c = ws[f"{col}{r}"]
            c.value = f[col]
            c._style = copy(body_src._style)
            if numfmt:
                c.number_format = numfmt
    ws.column_dimensions.group("AN", "AP", hidden=True)

    table = ws.tables["MintresponseWorkOrders"]
    table.ref = f"A10:AP{LAST}"
    table.autoFilter.ref = table.ref
    next_id = max(tc.id for tc in table.tableColumns) + 1
    for i, (_col, label, _w, _f) in enumerate(HELPERS):
        table.tableColumns.append(TableColumn(id=next_id + i, name=label))

    # Overdue rows: red text in the aging columns.
    ws.conditional_formatting.add(
        f"AL{FIRST}:AM{LAST}",
        FormulaRule(formula=[f"N($AM{FIRST})>0"], font=Font(color="B42318", bold=True)))


# ------------------------------------------------------------------ Clients

def build_clients(ws, n=200):
    title(ws, "Clients", "Totals per client from the Work Orders sheet. "
                         "Clients appear in the order they were first entered.")
    labels = ["#", "Client", "Work orders", "Invoice revenue", "Gross profit",
              "Net profit", "Balance due", "Overdue", "Revenue rank"]
    header_row(ws, 5, 1, labels)
    widths = [6, 32, 13, 17, 16, 16, 16, 16, 9]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    first, last = 6, 5 + n
    for k in range(1, n + 1):
        r = 5 + k
        ws[f"A{r}"] = k
        ws[f"B{r}"] = f'=IFERROR(INDEX({rng("B")},MATCH(A{r},{rng("AP")},0)),"")'
        ws[f"C{r}"] = f'=IF(B{r}="","",COUNTIFS({rng("B")},B{r},{rng("A")},"<>"))'
        for col, src in (("D", "Z"), ("E", "J"), ("F", "AE"), ("G", "AH")):
            ws[f"{col}{r}"] = f'=IF(B{r}="","",SUMIFS({rng(src)},{rng("B")},B{r}))'
        ws[f"H{r}"] = f'=IF(B{r}="","",SUMIFS({rng("AH")},{rng("B")},B{r},{rng("AM")},">0"))'
        ws[f"I{r}"] = (f'=IF(B{r}="","",COUNTIF(D${first}:D${last},">"&D{r})'
                       f'+COUNTIF(D${first}:D{r},D{r}))')
        for col in "DEFGH":
            ws[f"{col}{r}"].number_format = MONEY
        for col in "ABCDEFGHI":
            ws[f"{col}{r}"].border = LINE
        ws[f"A{r}"].font = Font(color=MUTED)
        ws[f"I{r}"].font = Font(color=MUTED)
    ws["B4"] = "Total"
    ws["B4"].font = Font(bold=True)
    for col in "CDEFGH":
        ws[f"{col}4"] = f"=SUM({col}{first}:{col}{last})"
        ws[f"{col}4"].font = Font(bold=True)
        ws[f"{col}4"].number_format = "0" if col == "C" else MONEY
    ws.freeze_panes = "C6"
    protect(ws)
    return first, last


# ------------------------------------------------------------------ Receivables

RECV_COLS = [  # (header, Work Orders column, number format, width)
    ("WO #", "A", "General", 14), ("Client", "B", "General", 22),
    ("Store name", "D", "General", 22), ("Invoice #", "Q", "General", 13),
    ("Invoice status", "V", "General", 15), ("Invoice sent", "U", DATE, 12),
    ("Payment due", "X", DATE, 12), ("Invoice revenue", "Z", MONEY, 15),
    ("Received", "AG", MONEY, 14), ("Balance due", "AH", MONEY, 15),
    ("Days overdue", "AM", "0", 12), ("Aging", "AL", "General", 15),
    ("Handler", "N", "General", 14),
]


def build_receivables(ws):
    title(ws, "Receivables", "Open balances on sent invoices, most overdue first. "
                             "Updates automatically from the Work Orders sheet.", col="C")
    first, last = 7, 6 + CAPACITY
    ws["C4"], ws["F4"], ws["I4"] = "Open balance", "Overdue", "Invoices open"
    ws["D4"] = f"=SUM(L{first}:L{last})"
    ws["G4"] = f'=SUMIF(M{first}:M{last},">0",L{first}:L{last})'
    ws["J4"] = f'=COUNT(L{first}:L{last})'
    for a, b, f in (("C4", "D4", MONEY), ("F4", "G4", MONEY), ("I4", "J4", "0")):
        ws[a].font = Font(bold=True, color=MUTED)
        ws[b].font = Font(bold=True, size=13, color=INK)
        ws[b].number_format = f
    header_row(ws, 6, 1, ["#", "Row"] + [h for h, *_ in RECV_COLS])
    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 5
    ws.column_dimensions["B"].hidden = True
    for i, (_h, _src, _f, w) in enumerate(RECV_COLS):
        ws.column_dimensions[get_column_letter(3 + i)].width = w
    for n in range(1, CAPACITY + 1):
        r = 6 + n
        ws[f"A{r}"] = n
        ws[f"A{r}"].font = Font(color=MUTED)
        ws[f"B{r}"] = f'=IFERROR(MATCH(A{r},{rng("AO")},0),"")'
        for i, (_h, src, numfmt, _w) in enumerate(RECV_COLS):
            c = ws.cell(row=r, column=3 + i)
            col = rng(src)
            c.value = f'=IF($B{r}="","",IF(INDEX({col},$B{r})="","",INDEX({col},$B{r})))'
            c.number_format = numfmt
            c.border = LINE
    ws.conditional_formatting.add(
        f"C{first}:O{last}",
        FormulaRule(formula=[f"N($M{first})>0"], font=Font(color="B42318")))
    ws.freeze_panes = "D7"
    protect(ws)


# ------------------------------------------------------------------ Work Order Card

def build_card(ws):
    title(ws, "Work Order Card", "Choose a WO # in the yellow box to see the whole job. "
                                 "Print with Ctrl+P.")
    for col, w in (("B", 24), ("C", 34), ("D", 3), ("E", 24), ("F", 34), ("H", 6)):
        ws.column_dimensions[col].width = w
    ws["B5"] = "WO #"
    ws["B5"].font = Font(bold=True, size=12, color=NAVY)
    pick = ws["C5"]
    pick.fill, pick.border = INPUT_FILL, BOX
    pick.font = Font(bold=True, size=12)
    pick.protection = Protection(locked=False)
    dv = DataValidation(type="list", formula1=f"={rng('A')}", allow_blank=True,
                        showErrorMessage=False)
    ws.add_data_validation(dv)
    dv.add("C5")
    ws["H5"] = f'=IF(C5="","",IFERROR(MATCH(C5,{rng("A")},0),""))'
    ws.column_dimensions["H"].hidden = True
    ws["E5"] = '=IF(C5="","◄ Choose or type a WO #",IF(H5="","WO # not found",""))'
    ws["E5"].font = Font(italic=True, color="B42318")

    def value(src_col):
        col = rng(src_col)
        return f'=IF($H$5="","",IF(INDEX({col},$H$5)="","",INDEX({col},$H$5)))'

    sections = list(FORM_SECTIONS) + [("Receivables", ["aging_bucket", "days_overdue"])]
    extra = {"aging_bucket": ("Aging bucket", "AL", "General"),
             "days_overdue": ("Days overdue", "AM", "0")}
    left = sections[:3] + [sections[-1]]
    right = sections[3:-1]
    for col_l, col_v, group in (("B", "C", left), ("E", "F", right)):
        r = 7
        for name, keys in group:
            ws.merge_cells(f"{col_l}{r}:{col_v}{r}")
            h = ws[f"{col_l}{r}"]
            h.value, h.fill, h.font = name, HEAD_FILL, HEAD_FONT
            r += 1
            for key in keys:
                if key in extra:
                    label, src, numfmt = extra[key]
                    computed = True
                else:
                    f = BY_KEY[key]
                    label, src, numfmt, computed = f.label, f.col, fmt_for(f.kind), f.computed
                ws[f"{col_l}{r}"] = label
                ws[f"{col_l}{r}"].font = Font(color=MUTED)
                v = ws[f"{col_v}{r}"]
                v.value = value(src)
                v.number_format = numfmt
                v.font = Font(bold=True, color=INK)
                v.alignment = Alignment(horizontal="left", wrap_text=True, vertical="top")
                v.fill = CALC_FILL if computed else PatternFill()
                ws[f"{col_l}{r}"].border = LINE
                v.border = LINE
                r += 1
            r += 1
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = "B2:F48"
    protect(ws)


# ------------------------------------------------------------------ Dashboard

def bar_chart(ws, title_text, cats, vals, anchor, number_format="0", height=6.2):
    ch = BarChart()
    ch.type = "bar"
    ch.style = 2
    ch.title = None
    ch.legend = None
    ch.height, ch.width = height, 15.5
    ch.gapWidth = 60
    ch.add_data(vals, titles_from_data=False)
    ch.set_categories(cats)
    s = ch.series[0]
    s.graphicalProperties.solidFill = "3F8F4A"
    s.graphicalProperties.line.noFill = True
    s.dLbls = DataLabelList()
    s.dLbls.showVal = True
    s.dLbls.numFmt = number_format
    for attr in ("showSerName", "showCatName", "showLegendKey", "showPercent"):
        setattr(s.dLbls, attr, False)
    ch.x_axis.scaling.orientation = "maxMin"      # keep the table's top-to-bottom order
    ch.x_axis.delete = False
    ch.y_axis.delete = True
    ch.y_axis.majorGridlines = None
    ws.add_chart(ch, anchor)


def section(ws, row, text, note=None):
    ws[f"B{row}"] = text
    ws[f"B{row}"].font = Font(bold=True, size=12, color=NAVY)
    if note:
        ws[f"E{row}"] = note
        ws[f"E{row}"].font = Font(size=9, color=MUTED)


def build_dashboard(ws, clients_first, clients_last):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 2
    for i in range(2, 14):
        ws.column_dimensions[get_column_letter(i)].width = 13
    ws.column_dimensions["B"].width = 18
    ws.row_dimensions[2].height = 26
    ws["C2"] = "Mintresponse — Work Order Dashboard"
    ws["C2"].font = Font(bold=True, size=18, color=NAVY)
    ws["C3"] = ('="All amounts in USD, excluding sales tax  ·  As of "'
                '&TEXT(TODAY(),"mmmm d, yyyy")')
    ws["C3"].font = Font(size=10, color=MUTED)
    if os.path.exists(LOGO):
        img = Image(LOGO)
        img.width, img.height = 50, 47
        ws.add_image(img, "B2")

    wos = "'Work Orders'!"
    tiles = [
        ("Work orders", f"={wos}A5", "0",
         f'=COUNTIFS({rng("A")},"<>",{rng("R")},"<>Completed",{rng("R")},"<>Cancelled")&" open"'),
        ("Invoiced revenue", f"={wos}B5", MONEY, None),
        ("Gross profit*", f"={wos}C5", MONEY,
         f'=IF(N({wos}B5)=0,"",TEXT({wos}C5/{wos}B5,"0.0%")&" margin")'),
        ("Net profit*", f"={wos}D5", MONEY,
         f'=IF(N({wos}B5)=0,"",TEXT({wos}D5/{wos}B5,"0.0%")&" margin")'),
        ("Balance due*", f"={wos}E5", MONEY, f'=COUNTIF({rng("AL")},"?*")&" open invoices"'),
        ("Overdue", f'=SUMIFS({rng("AH")},{rng("AM")},">0")', MONEY,
         f'=COUNTIF({rng("AM")},">0")&" invoices past due"'),
        ("Needs financial input", f"={wos}F5", "0", '="work orders with an Enter… status"'),
        ("Over the approved NTE", f'=COUNTIF({rng("AJ")},"*exceeds NTE")', "0",
         '="quotes or invoices above NTE"'),
    ]
    for i, (label, formula, numfmt, sub) in enumerate(tiles):
        top = 5 if i < 4 else 9
        c0 = 2 + (i % 4) * 3
        c1 = c0 + 2
        warn = label in ("Overdue", "Needs financial input", "Over the approved NTE")
        for r in range(top, top + 3):
            ws.merge_cells(start_row=r, start_column=c0, end_row=r, end_column=c1)
            for c in range(c0, c1 + 1):
                cell = ws.cell(row=r, column=c)
                cell.fill = TILE_FILL
                cell.border = Border(left=THIN if c == c0 else None,
                                     right=THIN if c == c1 else None,
                                     top=THIN if r == top else None,
                                     bottom=THIN if r == top + 2 else None)
        lab = ws.cell(row=top, column=c0, value=label.upper())
        lab.font = Font(bold=True, size=9, color=MUTED)
        val = ws.cell(row=top + 1, column=c0, value=formula)
        val.font = Font(bold=True, size=18, color=INK)
        val.number_format = numfmt
        val.alignment = Alignment(horizontal="left")
        if sub:
            s = ws.cell(row=top + 2, column=c0, value=sub)
            s.font = Font(size=9, color=MUTED)
        ws.row_dimensions[top + 1].height = 28
        if warn:
            ws.conditional_formatting.add(
                f"{get_column_letter(c0)}{top}:{get_column_letter(c1)}{top + 2}",
                FormulaRule(formula=[f"N(${get_column_letter(c0)}${top + 1})>0"],
                            fill=WARN_FILL))
    ws["B12"] = ("*Totals include calculated rows only. Complete financial inputs before "
                 "relying on profit totals.")
    ws["B12"].font = Font(size=9, italic=True, color=MUTED)

    # Work orders by status
    r0 = 14
    section(ws, r0, "Work orders by status")
    header_row(ws, r0 + 1, 2, ["WO status", "Count"])
    statuses = WO_STATUSES + ["No status"]
    for i, s in enumerate(statuses):
        r = r0 + 2 + i
        ws[f"B{r}"] = s
        if s == "No status":
            ws[f"C{r}"] = f'=COUNTIFS({rng("A")},"<>",{rng("R")},"")'
        else:
            ws[f"C{r}"] = f'=COUNTIF({rng("R")},B{r})'
        ws[f"B{r}"].border = ws[f"C{r}"].border = LINE
    end = r0 + 1 + len(statuses)
    bar_chart(ws, "Work orders by status", Reference(ws, min_col=2, min_row=r0 + 2, max_row=end),
              Reference(ws, min_col=3, min_row=r0 + 2, max_row=end), f"E{r0}")

    # Receivables aging
    r0 = 28
    section(ws, r0, "Receivables aging")
    header_row(ws, r0 + 1, 2, ["Days past due", "Balance due", "Invoices"])
    for i, b in enumerate(AGING):
        r = r0 + 2 + i
        ws[f"B{r}"] = b
        ws[f"C{r}"] = f'=SUMIFS({rng("AH")},{rng("AL")},B{r})'
        ws[f"D{r}"] = f'=COUNTIF({rng("AL")},B{r})'
        ws[f"C{r}"].number_format = MONEY
        for col in "BCD":
            ws[f"{col}{r}"].border = LINE
    end = r0 + 1 + len(AGING)
    bar_chart(ws, "Receivables aging", Reference(ws, min_col=2, min_row=r0 + 2, max_row=end),
              Reference(ws, min_col=3, min_row=r0 + 2, max_row=end), f"F{r0}",
              number_format='"$"#,##0')
    ws.column_dimensions["E"].width = 13

    # Top clients
    r0 = 42
    section(ws, r0, "Top clients by invoice revenue")
    header_row(ws, r0 + 1, 2, ["Client", "Revenue"])
    rank = f"Clients!$I${clients_first}:$I${clients_last}"
    for k in range(1, 9):
        r = r0 + 1 + k
        ws[f"B{r}"] = (f'=IFERROR(IF(INDEX(Clients!$D${clients_first}:$D${clients_last},'
                       f'MATCH({k},{rank},0))>0,INDEX(Clients!$B${clients_first}:$B${clients_last},'
                       f'MATCH({k},{rank},0)),""),"")')
        ws[f"C{r}"] = (f'=IF(B{r}="","",INDEX(Clients!$D${clients_first}:$D${clients_last},'
                       f'MATCH({k},{rank},0)))')
        ws[f"C{r}"].number_format = MONEY
        ws[f"B{r}"].border = ws[f"C{r}"].border = LINE
    bar_chart(ws, "Top clients", Reference(ws, min_col=2, min_row=r0 + 2, max_row=r0 + 9),
              Reference(ws, min_col=3, min_row=r0 + 2, max_row=r0 + 9), f"E{r0}",
              number_format='"$"#,##0')

    # Most overdue invoices
    r0 = 56
    section(ws, r0, "Most overdue invoices", "Full list on the Receivables sheet")
    cols = [("WO #", "C", "General"), ("Client", "D", "General"),
            ("Balance due", "L", MONEY), ("Days overdue", "M", "0"),
            ("Payment due", "I", DATE), ("Invoice #", "F", "General"),
            ("Handler", "O", "General")]
    header_row(ws, r0 + 1, 2, [c[0] for c in cols])
    for k in range(1, 11):
        r = r0 + 1 + k
        src = 6 + k                                # Receivables row k
        for i, (_h, col, numfmt) in enumerate(cols):
            c = ws.cell(row=r, column=2 + i)
            c.value = f'=IF(N(Receivables!$M${src})>0,Receivables!{col}{src},"")'
            c.number_format = numfmt
            c.border = LINE
    ws[f"B{r0 + 12}"] = (f'=IF(COUNTIF({rng("AM")},">0")=0,"No overdue invoices.",'
                         f'IF(COUNTIF({rng("AM")},">0")>10,"+ "&(COUNTIF({rng("AM")},">0")-10)'
                         f'&" more on the Receivables sheet",""))')
    ws[f"B{r0 + 12}"].font = Font(size=9, italic=True, color=MUTED)

    # How to use
    r0 = 71
    section(ws, r0, "How to use this workbook")
    tips = [
        "1. Enter work orders on the Work Orders sheet, one per row from row 11. Fill the yellow "
        "cells; blue/green cells calculate by themselves.",
        "2. Use the drop-down lists for WO status, Invoice status and Payment method.",
        "3. When a payment comes in, add it to 'Amount received' and update the Invoice status.",
        "4. This Dashboard, Receivables and Clients refresh automatically. Work Order Card shows "
        "any single job — pick the WO # in its yellow box.",
        "5. Definitions and the worked example are on the Guide and Example sheet.",
    ]
    for i, t in enumerate(tips):
        ws[f"B{r0 + 1 + i}"] = t
        ws[f"B{r0 + 1 + i}"].font = Font(size=10, color=INK)
    ws.page_setup.orientation = "portrait"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    protect(ws)


# ------------------------------------------------------------------ Guide additions

def update_guide(ws):
    for row in ws.iter_rows(min_col=2, max_col=2):
        c = row[0]
        if isinstance(c.value, str):
            c.value = (c.value.replace("Top totals include all 500 rows",
                                       f"Top totals include all {CAPACITY:,} rows")
                       .replace("Beyond row 510", f"Beyond row {LAST}"))
    ws.sheet_view.showGridLines = False


# ------------------------------------------------------------------ build

def fill(ws, rows):
    if len(rows) > CAPACITY:
        raise ValueError(f"The workbook holds {CAPACITY} work orders; got {len(rows)}.")
    for i, wo in enumerate(rows):
        r = FIRST + i
        for f in INPUT_FIELDS:
            v = wo.get(f.key)
            if v is not None and f.kind == "date":
                v = date.fromisoformat(v)
            ws[f"{f.col}{r}"].value = v


def build(out_path, rows=None):
    wb = load_workbook(TEMPLATE)
    work = wb["Work Orders"]
    build_work_orders(work)
    if rows:
        fill(work, rows)

    clients = wb.create_sheet("Clients")
    c_first, c_last = build_clients(clients)
    recv = wb.create_sheet("Receivables")
    build_receivables(recv)
    card = wb.create_sheet("Work Order Card")
    build_card(card)
    if rows:
        card["C5"] = rows[0]["wo_number"]
    dash = wb.create_sheet("Dashboard", 0)
    build_dashboard(dash, c_first, c_last)
    update_guide(wb["Guide and Example"])

    order = ["Dashboard", "Work Orders", "Work Order Card", "Receivables", "Clients",
             "Guide and Example"]
    wb._sheets = [wb[n] for n in order]
    colors = {"Dashboard": GREEN, "Work Orders": "E0B84B", "Work Order Card": NAVY,
              "Receivables": "B42318", "Clients": NAVY, "Guide and Example": "8A96A3"}
    for name, color in colors.items():
        wb[name].sheet_properties.tabColor = color
    wb.active = 0
    for ws in wb.worksheets:
        ws.sheet_view.tabSelected = ws.title == "Dashboard"
    wb.calculation.fullCalcOnLoad = True
    wb.save(out_path)
    return out_path


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("out", nargs="?", default=os.path.join(HERE, "Mintresponse_Work_Order_App.xlsx"))
    ap.add_argument("--demo", action="store_true", help="include 40 sample work orders")
    a = ap.parse_args()
    rows = None
    if a.demo:
        from seed_demo import demo_rows
        rows = demo_rows()
    print("Wrote", build(a.out, rows))


if __name__ == "__main__":
    main()
